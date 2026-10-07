import os
import json
import asyncio
import anthropic
import docx
import fitz
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Font
from config import *

PROFILE_DIR = "C:\\jobagent\\profile\\"
EXCEL_FILE = "C:\\jobagent\\vacancies.xlsx"
PENDING_FILE = "C:\\jobagent\\pending_jobs.json"

def read_docx(path: str) -> str:
    try:
        doc = docx.Document(path)
        return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
    except:
        return ""

def read_pdf(path: str) -> str:
    try:
        doc = fitz.open(path)
        text = ""
        for page in doc:
            text += page.get_text()
        return text[:3000]
    except:
        return ""

def load_profile() -> str:
    """Читает все файлы профиля и собирает в один текст"""
    profile_text = ""
    priority = ["CV_EN_v2.docx", "CV_RU_v2.docx", "CV_HeadOfDept.docx"]

    for filename in os.listdir(PROFILE_DIR):
        path = os.path.join(PROFILE_DIR, filename)
        if filename.endswith(".docx"):
            text = read_docx(path)
        elif filename.endswith(".pdf") and "photo" not in filename:
            text = read_pdf(path)
        else:
            continue
        if text:
            profile_text += f"\n=== {filename} ===\n{text[:1500]}\n"

    return profile_text[:8000]

def init_excel():
    """Создаёт Excel файл если не существует"""
    if os.path.exists(EXCEL_FILE):
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Вакансии"
    headers = ["#", "Дата", "Должность", "Компания", "Источник",
               "Ссылка", "Оценка", "Комментарий", "Статус"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="1C1C2E")
        cell.font = Font(bold=True, color="FFFFFF")
    wb.save(EXCEL_FILE)

def save_to_excel(job: dict, score: int, comment: str):
    """Сохраняет вакансию в Excel"""
    wb = load_workbook(EXCEL_FILE)
    ws = wb.active
    row = ws.max_row + 1
    num = row - 1

    ws.cell(row=row, column=1, value=num)
    ws.cell(row=row, column=2, value=datetime.now().strftime("%d.%m.%Y"))
    ws.cell(row=row, column=3, value=job.get("title", ""))
    ws.cell(row=row, column=4, value=job.get("company", ""))
    ws.cell(row=row, column=5, value=job.get("source", ""))
    ws.cell(row=row, column=6, value=job.get("link", ""))
    ws.cell(row=row, column=7, value=score)
    ws.cell(row=row, column=8, value=comment)

    status_cell = ws.cell(row=row, column=9, value="Нет ответа")
    status_cell.fill = PatternFill("solid", fgColor="FFD700")

    if score >= 8:
        ws.cell(row=row, column=7).fill = PatternFill("solid", fgColor="00C851")
    elif score >= 6:
        ws.cell(row=row, column=7).fill = PatternFill("solid", fgColor="FFD700")
    else:
        ws.cell(row=row, column=7).fill = PatternFill("solid", fgColor="FF4444")

    wb.save(EXCEL_FILE)
    return num

def analyze_match(job: dict, profile: str) -> tuple:
    """Claude оценивает соответствие вакансии профилю"""
    client = anthropic.Anthropic(api_key=ANTHROPIC_KEY)
    try:
        message = client.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=500,
            messages=[{"role": "user", "content": f"""Оцени насколько эта вакансия подходит кандидату на позиции Senior Project Manager или Program Manager.

ПРОФИЛЬ КАНДИДАТА:
{profile[:4000]}

ВАКАНСИЯ:
Должность: {job.get('title')}
Компания: {job.get('company')}
Почему найдена: {job.get('reason', '')}
Ссылка: {job.get('link')}

Сопоставь обязанности и уровень вакансии с опытом кандидата в управлении
проектами и программами. Не приписывай кандидату отсутствующие в профиле
навыки, опыт или достижения.

Верни JSON:
{{"score": 1-10, "comment": "краткий комментарий почему подходит или нет (2-3 предложения)"}}

Только JSON."""}]
        )
        result = message.content[0].text.strip()
        result = result.replace("```json","").replace("```","").strip()
        data = json.loads(result)
        return data.get("score", 5), data.get("comment", "")
    except Exception as e:
        return 5, f"Ошибка анализа: {e}"

async def run_profile_agent():
    """Читает pending вакансии и анализирует их"""
    print(f"\n[{datetime.now().strftime('%H:%M:%S')}] Profile агент запущен")

    init_excel()

    if not os.path.exists(PENDING_FILE):
        print("  Нет вакансий для анализа")
        return

    with open(PENDING_FILE, "r", encoding="utf-8") as f:
        pending = json.load(f)

    if not pending:
        print("  Нет вакансий для анализа")
        return

    print(f"  Загружаю профиль из {PROFILE_DIR}...")
    profile = load_profile()
    print(f"  Профиль загружен: {len(profile)} символов")
    print(f"  Анализирую {len(pending)} вакансий...")

    for job_id, job in pending.items():
        score, comment = analyze_match(job, profile)
        num = save_to_excel(job, score, comment)
        print(f"  [{num}] {job.get('title')} — оценка: {score}/10")

    print(f"  Сохранено в {EXCEL_FILE}")

def job():
    asyncio.run(run_profile_agent())

job()